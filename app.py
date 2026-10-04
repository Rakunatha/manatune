"""Manatune: sign in with Google, use three AI tools, track activity, file a provenance report with an IP advocate.

Creator flow:  /login -> AI tools -> Track & File (report) -> send to an IP advocate.
Advocate flow: /ipadvo (separate login) -> accept / evaluate / file -> download .docx.
"""
import os, io, re, json, time, hashlib, uuid, secrets, zipfile
from datetime import datetime
from functools import wraps
from urllib.parse import urlsplit, urlunsplit
import requests
from flask import Flask, request, jsonify, render_template, redirect, flash, url_for, session, g, send_file
from flask_sqlalchemy import SQLAlchemy
from flask_login import LoginManager, UserMixin, login_user, logout_user, login_required, current_user
from jinja2 import DictLoader
from sqlalchemy import UniqueConstraint
from sqlalchemy.exc import OperationalError, IntegrityError
from authlib.integrations.flask_client import OAuth
from werkzeug.middleware.proxy_fix import ProxyFix

# The only built-in AI tools. The extension tracks exactly these three domains.
TOOLS = [
    {"key": "rdxper", "name": "rdxper.space", "domain": "rdxper.space", "url": "https://rdxper.space"},
    {"key": "fashiqai", "name": "fashiqai.com", "domain": "fashiqai.com", "url": "https://fashiqai.com"},
    {"key": "dratido", "name": "dratido.onrender.com", "domain": "dratido.onrender.com", "url": "https://dratido.onrender.com"},
]
TOOL = {t["key"]: t for t in TOOLS}
PROTECTION = {"copyright": "Copyright", "trademark": "Trademark", "design": "Design", "other": "Other (advocate to advise)"}
STATUS = {"submitted": "Sent, waiting for the advocate", "accepted": "Accepted", "filed": "Filed", "declined": "Declined"}
NOTE = "Evidence only: reported by the Manatune browser extension, not proof of authorship or ownership, and not legal advice."

# ======================= pages =======================
TEMPLATES = {
    "login.html": r'''<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Manatune: sign in</title>
<style>
:root{--bg:#f3f5f9;--card:#fff;--text:#11151b;--mut:#5b6471;--bd:#d9dee7;--blue:#1f7ae8;--blue2:#1665c4;--bad:#c0392b}
@media (prefers-color-scheme:dark){:root{--bg:#0f1114;--card:#181b20;--text:#fff;--mut:#a7adb6;--bd:#262b33;--blue:#2d8cff;--blue2:#1a5fb4;--bad:#ff7a6b}}
*{box-sizing:border-box}
body{margin:0;min-height:100vh;display:grid;place-items:center;background:var(--bg);color:var(--text);font:14px/1.45 -apple-system,BlinkMacSystemFont,"Segoe UI",Helvetica,Arial,sans-serif;padding:16px}
.box{width:100%;max-width:380px}
.logo{color:var(--blue);font-size:30px;font-weight:800;letter-spacing:-.02em;margin-bottom:18px}
.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:20px}
label{display:block;font-size:12px;color:var(--mut);margin:12px 0 4px}
input,select{width:100%;background:var(--bg);color:var(--text);border:1px solid var(--bd);border-radius:6px;padding:10px;font:inherit}
input:focus,select:focus,button:focus-visible{outline:2px solid var(--blue);outline-offset:1px}
.go{width:100%;margin-top:18px;background:var(--blue);color:#fff;border:0;border-radius:20px;padding:11px;font-weight:700;font-size:14px;cursor:pointer}
.go:hover{background:var(--blue2)}
.err{color:var(--bad);margin:0 0 4px}
.hint{color:var(--mut);font-size:12px;margin:6px 0 0}
.g{display:block;text-align:center;text-decoration:none;color:var(--text);background:var(--bg);border:1px solid var(--bd);border-radius:20px;padding:10px;font-weight:700}
.g:hover{border-color:var(--blue)}
</style>
</head>
<body>
<div class="box">
  <div class="logo">Manatune</div>
  <div class="card">
    {% for m in get_flashed_messages() %}<p class="err">{{ m }}</p>{% endfor %}
    {% if mode == 'advocate' %}
    <p style="margin:0 0 12px"><b>IP advocate sign-in</b></p>
    <p class="hint" style="margin:0 0 14px">For lawyers and firms who receive filing requests from creators.</p>
    <a class="g" href="/auth/google/advocate">Continue with Google</a>
    <p class="hint" style="text-align:center;margin-top:12px"><a href="/login" style="color:var(--mut)">Not an advocate? Creator sign-in</a></p>
    {% else %}
    <a class="g" href="/auth/google">Continue with Google</a>
    <p class="hint" style="text-align:center;margin-top:12px"><a href="/ipadvo" style="color:var(--mut)">IP advocate or firm? Sign in here</a></p>
    {% endif %}
  </div>
</div>
</body>
</html>
''',
    "extension.html": r'''<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Manatune: extension</title>
<style>
:root{--bg:#f3f5f9;--card:#fff;--text:#11151b;--mut:#5b6471;--bd:#d9dee7;--blue:#1f7ae8;--bad:#c0392b}
@media (prefers-color-scheme:dark){:root{--bg:#0f1114;--card:#181b20;--text:#fff;--mut:#a7adb6;--bd:#262b33;--blue:#2d8cff;--bad:#ff7a6b}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:14px/1.45 -apple-system,Segoe UI,Arial,sans-serif}
.w{max-width:680px;margin:0 auto;padding:20px 16px}.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:16px;margin-bottom:14px}
h1{font-size:22px;margin:0 0 6px}h2{font-size:16px;margin:0 0 6px}.sub{color:var(--mut)}
.btn{display:inline-block;background:var(--blue);color:#fff;border:0;border-radius:20px;padding:8px 16px;font:inherit;font-weight:700;cursor:pointer;text-decoration:none}
.btn.ghost{background:none;color:var(--blue);border:1px solid var(--bd);padding:4px 12px;font-size:12px}
.mono{font-family:ui-monospace,Menlo,monospace;font-size:12px;word-break:break-all;color:var(--mut)}.err{color:var(--bad)}
.row{display:flex;justify-content:space-between;gap:8px;align-items:center;padding:8px 0;border-top:1px solid var(--bd)}
a{color:var(--blue)}
</style></head><body><div class="w">
<p><a href="/">&larr; Back</a></p><h1>Browser extension</h1>
<p class="sub">It records which pages you open and what you download on the three AI tools. Metadata only, never page content.</p>
<div class="card"><h2>1. Install</h2><ol><li>Download and unzip.</li><li>Open <b>chrome://extensions</b>, switch on <b>Developer mode</b>.</li><li><b>Load unpacked</b> and pick the folder.</li></ol>
<a class="btn" href="/extension/download.zip">Download extension</a></div>
<div class="card"><h2>2. Connect</h2><p class="sub">Generate a token and paste it into the extension.</p>
<button class="btn" id="gen">Generate token</button><div id="fresh" class="mono" style="margin-top:10px"></div><p class="err" id="terr"></p><div id="toks"></div></div>
</div>
<script>
const $=s=>document.querySelector(s);
const esc=s=>String(s==null?'':s).replace(/[&<>"']/g,m=>({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[m]));
async function api(p,o){const r=await fetch(p,Object.assign({headers:{'Content-Type':'application/json'}},o));const d=await r.json().catch(()=>({}));if(!r.ok)throw new Error(d.error||'Something went wrong');return d}
async function load(){const t=await api('/api/tokens');$('#toks').innerHTML=t.map(x=>'<div class="row"><span class="mono">'+esc(x.prefix)+'... created '+esc(x.created.slice(0,10))+'</span><button class="btn ghost" data-r="'+x.id+'">Revoke</button></div>').join('')}
$('#gen').onclick=async()=>{$('#terr').textContent='';try{const d=await api('/api/tokens',{method:'POST',body:'{}'});$('#fresh').textContent='Copy this now, it is shown once: '+d.token;load()}catch(e){$('#terr').textContent=e.message}};
$('#toks').onclick=async e=>{const id=e.target.dataset&&e.target.dataset.r;if(!id)return;await api('/api/tokens/'+id,{method:'DELETE'});$('#fresh').textContent='';load()};
load().catch(e=>{$('#terr').textContent=e.message});
</script></body></html>
''',
    "workspace.html": r'''<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Manatune</title>
<style>
:root{--bg:#f3f5f9;--card:#fff;--text:#11151b;--mut:#5b6471;--bd:#d9dee7;--blue:#1f7ae8;--blue2:#1665c4;--on:#e6f0fd;--ok:#168246;--bad:#c0392b}
@media (prefers-color-scheme:dark){:root{--bg:#0f1114;--card:#181b20;--text:#fff;--mut:#a7adb6;--bd:#262b33;--blue:#2d8cff;--blue2:#1a5fb4;--on:#14283f;--ok:#4cc282;--bad:#ff7a6b}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:14px/1.45 -apple-system,Segoe UI,Arial,sans-serif}
button,input,textarea,select{font:inherit;color:inherit}:focus-visible{outline:2px solid var(--blue);outline-offset:2px}
header{display:flex;justify-content:space-between;align-items:center;gap:10px;flex-wrap:wrap;padding:10px 16px;background:var(--card);border-bottom:1px solid var(--bd)}
.brand{color:var(--blue);font-size:26px;font-weight:800;letter-spacing:-.02em}
.top{display:flex;gap:6px;align-items:center;flex-wrap:wrap}.top a,.top button{color:var(--mut);background:none;border:0;text-decoration:none;padding:6px 8px;cursor:pointer}
.tabs{display:flex;gap:4px;padding:0 16px;background:var(--card);border-bottom:1px solid var(--bd)}
.tabs button{background:none;border:0;border-bottom:3px solid transparent;padding:10px 14px;font-weight:700;color:var(--mut);cursor:pointer}
.tabs button[aria-current=page]{color:var(--text);border-bottom-color:var(--blue)}
main{max-width:900px;margin:0 auto;padding:20px 16px 80px}
h1{font-size:22px;margin:0 0 4px}.sub,.sm{color:var(--mut)}.sm{font-size:12px}
.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px;margin-bottom:10px}
.sp{display:flex;justify-content:space-between;gap:10px;align-items:center;flex-wrap:wrap}
.btn,.use{display:inline-block;background:var(--blue);color:#fff;border:0;border-radius:20px;padding:8px 16px;font-weight:700;font-size:13px;cursor:pointer;text-decoration:none}
.btn:hover,.use:hover{background:var(--blue2)}.btn:disabled{opacity:.6}.btn.ghost{background:none;color:var(--blue);border:1px solid var(--bd)}
.pill{display:inline-block;border-radius:12px;padding:2px 10px;font-size:12px;font-weight:700;background:var(--on)}
.mono{font-family:ui-monospace,Menlo,monospace;font-size:12px;word-break:break-all;color:var(--mut)}
.note{background:var(--on);border-radius:10px;padding:10px 14px;margin:0 0 14px}.note.bad{color:var(--bad)}
label{display:block;font-size:12px;color:var(--mut);margin:12px 0 4px}
input,textarea,select{width:100%;background:var(--bg);border:1px solid var(--bd);border-radius:6px;padding:10px}
h2{font-size:16px;margin:24px 0 10px}
</style></head><body>
<header><span class="brand">Manatune</span><div class="top"><a href="/extension">Extension</a><span class="sm">{{ me.name }}</span>
<form method="post" action="/logout" style="margin:0"><button>Sign out</button></form></div></header>
<nav class="tabs" id="tabs"></nav>
<main id="view" aria-live="polite"></main>
<script>
const $=s=>document.querySelector(s),view=$("#view");
const el=(t,c,x)=>{const e=document.createElement(t);if(c)e.className=c;if(x!=null)e.textContent=x;return e};
const fld=(l,n)=>{const b=el("label",null,l);b.append(n);return b};
const safe=u=>/^https?:\/\//i.test(u||"")?u:"#";const d10=s=>(s||"").slice(0,10);
const api=async(u,b,m)=>{const o=b===undefined?{}:{method:m||"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(b)};
  const r=await fetch(u,o);if(r.status===401){location.href="/login";throw 0}
  const d=await r.json().catch(()=>({}));if(!r.ok)throw new Error(d.error||"Request failed");return d};
function fail(e){if(e===0)return;view.querySelectorAll(".note.bad").forEach(n=>n.remove());view.prepend(el("div","note bad",e.message||"Something went wrong."))}
async function act(b,fn){b.disabled=true;try{await fn()}catch(e){fail(e)}b.disabled=false}
function head(t,s){view.replaceChildren(el("h1",null,t),el("p","sub",s))}
const ST={{ status|tojson }};
let tab={{ tab|tojson }};
const TABS=[["tools","AI tools"],["track","Track & File"]];
function show(t){tab=t;$("#tabs").replaceChildren(...TABS.map(([k,l])=>{const b=el("button",null,l);if(k===t)b.setAttribute("aria-current","page");b.onclick=()=>show(k);return b}));(t==="tools"?vTools:vTrack)()}

async function vTools(){head("AI tools","Open a tool from here so your activity and downloads are tracked.");
  let ts;try{ts=await api("/api/tools")}catch(e){return fail(e)}
  ts.forEach(t=>{const c=el("div","card sp"),l=el("div");l.append(el("div",null,t.name),el("div","sm",t.url));
    const a=el("a","use","Open tool");a.href="/go/"+t.key;a.target="_blank";a.rel="noopener";c.append(l,a);view.append(c)});
  view.append(el("div","note","Tracking needs the browser extension (link at the top). Then everything you download from these tools shows up in Track & File."))}

async function vTrack(){head("Track & File","Your tracked activity, the provenance report for it, and the IP filing.");
  let sm,reps,fil;try{[sm,reps,fil]=await Promise.all([api("/api/summary"),api("/api/reports"),api("/api/filings")])}catch(e){return fail(e)}
  if(!sm.extension_active){const n=el("div","note");n.append("No browser activity has arrived yet. Install the ");const a=el("a",null,"extension");a.href="/extension";n.append(a,", then open a tool and download something.");view.append(n)}
  view.append(el("h2",null,"1. Tracked activity"));
  sm.tools.forEach(t=>{const c=el("div","card sp"),l=el("div");l.append(el("div",null,t.name),el("div","sm",t.visits+" visits, "+t.downloads+" downloads, last activity "+(d10(t.last)||"none yet")));
    const b=el("button","btn","Create report");b.disabled=!t.downloads;b.title=t.downloads?"":"Download something from this tool first";
    b.onclick=()=>act(b,async()=>{await api("/api/reports",{tool:t.key});vTrack()});c.append(l,b);view.append(c)});
  view.append(el("h2",null,"2. Provenance reports"));
  if(!reps.length)view.append(el("div","card sm","No reports yet."));
  const advs=reps.length?await api("/api/advocates").catch(()=>[]):[];
  reps.forEach(r=>{const c=el("div","card"),top=el("div","sp");top.append(el("b",null,r.title),el("span","pill",r.filing?ST[r.filing.status]:"Not filed"));
    c.append(top,el("div","sm",d10(r.period.start)+" to "+d10(r.period.end)));
    const body=el("div");r.narrative.split("\n\n").forEach(p=>body.append(el("p",null,p)));c.append(body);
    r.outputs.forEach(o=>{const w=el("div");w.style.cssText="padding:8px 0;border-top:1px solid var(--bd)";const a=el("a",null,o.filename);a.href=safe(o.link);a.target="_blank";a.rel="noopener noreferrer";
      w.append(a,el("div","sm",(o.mime||"type unknown")+" | "+(o.size==null?"size unknown":o.size+" bytes")+" | "+o.at),el("div","mono","SHA-256: "+(o.sha256||"not recorded")));c.append(w)});
    const acts=el("div","sp");acts.style.cssText="justify-content:flex-start;gap:8px;margin-top:8px";
    const dx=el("a","btn ghost","Download .docx");dx.href="/api/reports/"+r.id+"/docx";acts.append(dx);
    if(!r.filing||r.filing.status==="declined"){const ff=el("button","btn","File for IP protection"),f=el("div");f.style.display="none";
      const as=document.createElement("select");advs.forEach(a=>{const o=el("option",null,a.name);o.value=a.id;as.append(o)});
      const ty=document.createElement("select");Object.entries({copyright:"Copyright",trademark:"Trademark",design:"Design",other:"Other (advocate to advise)"}).forEach(([v,l])=>{const o=el("option",null,l);o.value=v;ty.append(o)});
      const ms=el("textarea");ms.rows=2;ms.maxLength=1000;ms.placeholder="What you did yourself, and anything else the advocate should know";
      const sd=el("button","btn","Send to advocate");sd.style.marginTop="12px";
      sd.onclick=()=>act(sd,async()=>{await api("/api/filings",{report_id:r.id,advocate_id:parseInt(as.value,10),protection_type:ty.value,message:ms.value});vTrack()});
      f.append(fld("IP advocate or firm",as),fld("Protection sought",ty),fld("Message (optional)",ms),sd);
      ff.onclick=()=>{if(!advs.length)return fail(new Error("No IP advocates have signed up yet."));f.style.display=f.style.display==="none"?"block":"none"};
      acts.append(ff);c.append(acts,f)}else c.append(acts);
    if(r.filing&&r.filing.note){const n=el("div","note","Advocate: "+r.filing.note);n.style.marginTop="8px";c.append(n)}
    if(r.filing&&r.filing.filing_ref)c.append(el("div","sm","Filing reference: "+r.filing.filing_ref));
    view.append(c)});
}
show(tab);
</script></body></html>
''',
    "advocate_home.html": r'''<!DOCTYPE html>
<html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Manatune: IP advocate portal</title>
<style>
:root{--bg:#f3f5f9;--card:#fff;--text:#11151b;--mut:#5b6471;--bd:#d9dee7;--blue:#1f7ae8;--blue2:#1665c4;--on:#e6f0fd;--bad:#c0392b}
@media (prefers-color-scheme:dark){:root{--bg:#0f1114;--card:#181b20;--text:#fff;--mut:#a7adb6;--bd:#262b33;--blue:#2d8cff;--blue2:#1a5fb4;--on:#14283f;--bad:#ff7a6b}}
*{box-sizing:border-box}body{margin:0;background:var(--bg);color:var(--text);font:14px/1.45 -apple-system,Segoe UI,Arial,sans-serif}
button,input,textarea{font:inherit;color:inherit}:focus-visible{outline:2px solid var(--blue);outline-offset:2px}
header{display:flex;justify-content:space-between;align-items:center;gap:10px;flex-wrap:wrap;padding:10px 16px;background:var(--card);border-bottom:1px solid var(--bd)}
.brand{color:var(--blue);font-size:24px;font-weight:800}.w{max-width:900px;margin:0 auto;padding:20px 16px 80px}
h1{font-size:22px;margin:0 0 4px}.sub,.sm{color:var(--mut)}.sm{font-size:12px}
.card{background:var(--card);border:1px solid var(--bd);border-radius:10px;padding:14px;margin-bottom:10px}
.sp{display:flex;justify-content:space-between;gap:10px;align-items:center;flex-wrap:wrap}
.btn{display:inline-block;background:var(--blue);color:#fff;border:0;border-radius:20px;padding:8px 16px;font-weight:700;font-size:13px;cursor:pointer;text-decoration:none}
.btn:hover{background:var(--blue2)}.btn:disabled{opacity:.6}.btn.ghost{background:none;color:var(--blue);border:1px solid var(--bd)}.btn.bad{background:var(--bad)}
.pill{display:inline-block;border-radius:12px;padding:2px 10px;font-size:12px;font-weight:700;background:var(--on)}
.mono{font-family:ui-monospace,Menlo,monospace;font-size:12px;word-break:break-all;color:var(--mut)}
.note{background:var(--on);border-radius:10px;padding:10px 14px;margin:0 0 14px}.note.bad{color:var(--bad)}
label{display:block;font-size:12px;color:var(--mut);margin:12px 0 4px}input,textarea{width:100%;background:var(--bg);border:1px solid var(--bd);border-radius:6px;padding:10px}
.row{cursor:pointer}.row:hover{border-color:var(--blue)}
</style></head><body>
<header><span class="brand">Manatune <span class="sm">IP advocate portal</span></span><span class="sp" style="gap:12px"><span class="sm">{{ me.name }}</span>
<form method="post" action="/logout" style="margin:0"><button class="btn ghost">Sign out</button></form></span></header>
<div class="w" id="view" aria-live="polite"></div>
<script>
const $=s=>document.querySelector(s),view=$("#view");
const el=(t,c,x)=>{const e=document.createElement(t);if(c)e.className=c;if(x!=null)e.textContent=x;return e};
const fld=(l,n)=>{const b=el("label",null,l);b.append(n);return b};
const safe=u=>/^https?:\/\//i.test(u||"")?u:"#";const d10=s=>(s||"").slice(0,10);
const ST={{ status|tojson }};
const api=async(u,b)=>{const o=b===undefined?{}:{method:"POST",headers:{"Content-Type":"application/json"},body:JSON.stringify(b)};
  const r=await fetch(u,o);if(r.status===401){location.href="/ipadvo";throw 0}
  const d=await r.json().catch(()=>({}));if(!r.ok)throw new Error(d.error||"Request failed");return d};
function fail(e){if(e===0)return;view.querySelectorAll(".note.bad").forEach(n=>n.remove());view.prepend(el("div","note bad",e.message||"Something went wrong."))}
async function act(b,fn){b.disabled=true;try{await fn()}catch(e){fail(e)}b.disabled=false}
async function list(){view.replaceChildren(el("h1",null,"Filing requests"),el("p","sub","Accept, evaluate and file what creators send you, then download the filing as .docx."));
  let rows;try{rows=await api("/api/advocate/requests")}catch(e){return fail(e)}
  if(!rows.length)view.append(el("div","card sm","Nothing yet."));
  rows.forEach(r=>{const c=el("div","card row"),t=el("div","sp");t.append(el("b",null,r.title),el("span","pill",ST[r.status]));
    c.append(t,el("div","sm",r.creator_name+" | "+r.protection+" | received "+d10(r.created)));c.onclick=()=>detail(r.id);view.append(c)})}
async function detail(id){let f;try{f=await api("/api/advocate/requests/"+id)}catch(e){return fail(e)}
  view.replaceChildren();const back=el("button","btn ghost","Back");back.onclick=list;view.append(back);
  const h=el("h1",null,f.title);h.style.marginTop="14px";view.append(h);
  const top=el("div","sp");top.append(el("span","pill",ST[f.status]));const dx=el("a","btn","Download filing .docx");dx.href="/api/advocate/requests/"+f.id+"/docx";top.append(dx);view.append(top);
  const c=el("div","card");c.style.marginTop="12px";
  c.append(el("div",null,"Creator: "+f.creator_name+" ("+f.creator_email+")"),el("div",null,"Protection sought: "+f.protection),el("div",null,"AI tool: "+f.tool+" | "+d10(f.period.start)+" to "+d10(f.period.end)),el("div",null,"Creator's message: "+(f.message||"none")));
  if(f.filing_ref)c.append(el("div",null,"Filing reference: "+f.filing_ref));view.append(c);
  const s=el("div","card");f.narrative.split("\n\n").forEach(p=>s.append(el("p",null,p)));
  f.outputs.forEach(o=>{const w=el("div");w.style.cssText="padding:8px 0;border-top:1px solid var(--bd)";const a=el("a",null,o.filename);a.href=safe(o.link);a.target="_blank";a.rel="noopener noreferrer";
    w.append(a,el("div","sm",(o.mime||"type unknown")+" | "+(o.size==null?"size unknown":o.size+" bytes")+" | "+o.at),el("div","mono","SHA-256: "+(o.sha256||"not recorded")));s.append(w)});
  s.append(el("div","sm",f.note));view.append(s);
  if(["filed","declined"].includes(f.status))return;
  const a=el("div","card"),note=el("textarea");note.rows=2;note.maxLength=1000;
  const run=(b,action,extra)=>b.onclick=()=>act(b,async()=>{await api("/api/advocate/requests/"+f.id+"/decision",Object.assign({action,note:note.value},extra&&extra()));detail(f.id)});
  a.append(fld("Message to the creator (needed to decline)",note));const bar=el("div","sp");bar.style.cssText="justify-content:flex-start;gap:8px;margin-top:10px";
  if(f.status==="submitted"){const b=el("button","btn","Accept");run(b,"accept");bar.append(b)}
  const dc=el("button","btn bad","Decline");run(dc,"decline");bar.append(dc);a.append(bar);view.append(a);
  if(f.status==="accepted"){const e=el("div","card"),ev=el("textarea");ev.rows=5;ev.maxLength=5000;ev.value=f.evaluation||"";e.append(fld("Your evaluation (private, included in your .docx only)",ev));
    const sv=el("button","btn ghost","Save evaluation");sv.style.marginTop="10px";run(sv,"evaluate",()=>({evaluation:ev.value}));
    const rf=el("input");rf.maxLength=120;rf.placeholder="Application or registration number";e.append(sv,fld("Filing reference",rf));
    const fb=el("button","btn","Mark as filed");run(fb,"file",()=>({filing_ref:rf.value}));e.append(fb);view.append(e)}}
list();
</script></body></html>
''',
}


# ======================= app + models =======================
def _db_url():
    url = os.environ.get("DATABASE_URL", "sqlite:///manatune.db")
    for prefix in ("postgres://", "postgresql://"):
        if url.startswith(prefix):
            return "postgresql+psycopg2://" + url[len(prefix):]
    return url


BASE = os.path.dirname(os.path.abspath(__file__))
app = Flask(__name__, root_path=BASE)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.jinja_loader = DictLoader(TEMPLATES)
app.config.update(SECRET_KEY=os.environ.get("SECRET_KEY", "dev-only-change-me"), SQLALCHEMY_DATABASE_URI=_db_url(),
                  SQLALCHEMY_ENGINE_OPTIONS={"pool_pre_ping": True, "pool_recycle": 280}, SESSION_COOKIE_HTTPONLY=True,
                  SESSION_COOKIE_SAMESITE="Lax", SESSION_COOKIE_SECURE=bool(os.environ.get("RENDER")))
db = SQLAlchemy(app)
lm = LoginManager(app)


class User(UserMixin, db.Model):
    __tablename__ = "users"
    id = db.Column(db.Integer, primary_key=True)
    email = db.Column(db.String(255), unique=True, nullable=False)
    name = db.Column(db.String(120), nullable=False)
    role = db.Column(db.String(20), default="creator")  # creator | advocate
    pw = db.Column(db.String(255), nullable=False)  # unused legacy column (NOT NULL in existing databases)


class ApiToken(db.Model):  # personal token for the extension; only a SHA-256 hash is stored
    __tablename__ = "api_tokens"
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), index=True, nullable=False)
    token_hash = db.Column(db.String(64), unique=True, nullable=False)
    prefix = db.Column(db.String(12))
    created = db.Column(db.DateTime, default=datetime.utcnow)
    expires = db.Column(db.DateTime)  # column kept from the earlier version; unused
    last_used = db.Column(db.DateTime)  # unused, kept so existing databases still match
    revoked = db.Column(db.Boolean, default=False)


class Activity(db.Model):  # one metadata event: launch (from Manatune), visit or download (from the extension)
    __tablename__ = "tf_activity"
    __table_args__ = (UniqueConstraint("user_id", "client_id", name="uq_tf_activity_client"),)
    id = db.Column(db.Integer, primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), index=True, nullable=False)
    tool = db.Column(db.String(20), index=True, nullable=False)
    kind = db.Column(db.String(10), nullable=False)  # launch | visit | download
    client_id = db.Column(db.String(64))
    url = db.Column(db.Text)
    page_url = db.Column(db.Text)
    title = db.Column(db.String(300))
    filename = db.Column(db.String(255))
    mime = db.Column(db.String(100))
    size = db.Column(db.BigInteger)
    sha256 = db.Column(db.String(64))
    created = db.Column(db.DateTime, default=datetime.utcnow, index=True)


class Report(db.Model):  # snapshot of one tool's activity at the moment it was created
    __tablename__ = "tf_reports"
    id = db.Column(db.String(32), primary_key=True)
    user_id = db.Column(db.Integer, db.ForeignKey("users.id"), index=True, nullable=False)
    tool = db.Column(db.String(20), nullable=False)
    title = db.Column(db.String(200), nullable=False)
    narrative = db.Column(db.Text, nullable=False)
    data = db.Column(db.Text, nullable=False)  # JSON: period, summary, outputs
    created = db.Column(db.DateTime, default=datetime.utcnow)


class Filing(db.Model):
    __tablename__ = "tf_filings"
    id = db.Column(db.String(32), primary_key=True)
    report_id = db.Column(db.String(32), db.ForeignKey("tf_reports.id"), index=True, nullable=False)
    creator_id = db.Column(db.Integer, db.ForeignKey("users.id"), index=True, nullable=False)
    advocate_id = db.Column(db.Integer, db.ForeignKey("users.id"), index=True, nullable=False)
    protection_type = db.Column(db.String(20), nullable=False)
    message = db.Column(db.String(1000))
    status = db.Column(db.String(20), default="submitted", index=True)  # submitted | accepted | filed | declined
    evaluation = db.Column(db.Text)  # private to the advocate
    note = db.Column(db.String(1000))  # advocate's message to the creator
    filing_ref = db.Column(db.String(120))
    created = db.Column(db.DateTime, default=datetime.utcnow)


# ======================= helpers =======================
def err(m, c):
    return jsonify(error=m), c


def iso(d):
    return d.isoformat(timespec="seconds") + "Z"


def now():
    return datetime.utcnow().replace(microsecond=0)


def clean(v, n):
    v = " ".join(str(v).split()) if v is not None else ""
    return v[:n] or None


def web_url(v, n=2000):
    """A normalised http(s) URL without fragment or credentials, or None."""
    try:
        s = urlsplit(str(v or "").strip())
        if s.scheme not in ("http", "https") or not s.hostname or s.username or s.password or len(str(v)) > n:
            return None
        return urlunsplit((s.scheme, s.netloc, s.path or "/", s.query, ""))
    except ValueError:
        return None


def tool_for_url(u):
    try:
        s = urlsplit(str(u or ""))
        host = (s.hostname or "").lower()
        if s.scheme not in ("http", "https") or not host:
            return None
    except ValueError:
        return None
    return next((t for t in TOOLS if host == t["domain"] or host.endswith("." + t["domain"])), None)


def role_api(role):
    def deco(f):
        @wraps(f)
        def w(*a, **k):
            if not current_user.is_authenticated:
                return err("login required", 401)
            if current_user.role != role:
                return err("This is for %s accounts." % role, 403)
            return f(*a, **k)
        return w
    return deco


creator_only, advocate_only = role_api("creator"), role_api("advocate")


@app.before_request
def same_site_only():  # extra CSRF guard on top of SameSite=Lax cookies
    if request.method not in ("GET", "HEAD", "OPTIONS") and request.headers.get("Sec-Fetch-Site") == "cross-site":
        return err("Cross-site requests are not allowed.", 403)


def token_required(f):  # extension requests: Bearer token, no cookies
    @wraps(f)
    def w(*a, **k):
        h = request.headers.get("Authorization", "")
        raw = h[7:].strip() if h.startswith("Bearer ") else ""
        t = ApiToken.query.filter_by(token_hash=hashlib.sha256(raw.encode()).hexdigest(), revoked=False).first() if raw.startswith("mt_") else None
        u = db.session.get(User, t.user_id) if t else None
        if not u or u.role != "creator":
            return err("Invalid token. Generate a new one on the Manatune website.", 401)
        g.ext_user = u
        return f(*a, **k)
    return w


# ======================= AI summary (optional Groq; plain text otherwise) =======================
SUMMARY_PROMPT = """Write the factual summary of a provenance report about someone's use of an AI tool.
Use ONLY the supplied JSON: who used which tool, when, how active they were, which files they downloaded.
Do not claim the person authored or owns the files, and do not mention prompts (they are not recorded).
Plain prose, 2 short paragraphs, no markdown. Return JSON: {"summary": str}"""


def plain_summary(snap):
    s, who, tool = snap["summary"], snap["creator"], snap["tool"]["name"]
    names = ", ".join("%s (%s)" % (o["filename"], o["mime"] or "type unknown") for o in snap["outputs"][:10])
    more = "" if len(snap["outputs"]) <= 10 else " and %d more" % (len(snap["outputs"]) - 10)
    return ("%s used %s between %s and %s (UTC). The browser extension recorded %d page visit(s) and %d download(s), "
            "about %d minute(s) of activity (an estimate).\n\nDownloaded files: %s%s." %
            (who, tool, snap["period"]["start"], snap["period"]["end"], s["visits"], s["downloads"], s["active_minutes"], names, more))


def summarise(snap):
    key = os.environ.get("GROQ_API_KEY")
    if key:
        try:
            r = requests.post("https://api.groq.com/openai/v1/chat/completions", timeout=30, headers={"Authorization": "Bearer " + key},
                              json={"model": os.environ.get("GROQ_MODEL", "llama-3.3-70b-versatile"), "temperature": 0.2,
                                    "response_format": {"type": "json_object"},
                                    "messages": [{"role": "system", "content": SUMMARY_PROMPT},
                                                 {"role": "user", "content": json.dumps({k: snap[k] for k in ("creator", "tool", "period", "summary", "outputs")})}]})
            r.raise_for_status()
            text = str(json.loads(r.json()["choices"][0]["message"]["content"]).get("summary", "")).strip()
            if 40 <= len(text) <= 3000:
                return text
        except Exception:
            app.logger.exception("AI summary failed; using the plain summary")
    return plain_summary(snap)


def snapshot(user, tool, evs):
    stamps = sorted(e.created for e in evs)
    gaps = [(b - a).total_seconds() for a, b in zip(stamps, stamps[1:])]
    minutes = int(round(sum(x for x in gaps if x <= 600) / 60.0))
    outs = [dict(filename=e.filename or "download", mime=e.mime, size=e.size, sha256=e.sha256, link=e.url or e.page_url, at=iso(e.created))
            for e in evs if e.kind == "download"][:200]
    return {"creator": user.name, "tool": {"name": tool["name"], "url": tool["url"]}, "period": {"start": iso(stamps[0]), "end": iso(stamps[-1])},
            "summary": {"visits": sum(e.kind == "visit" for e in evs), "downloads": len(outs), "active_minutes": minutes}, "outputs": outs}


def report_json(r, filing=None):
    d = json.loads(r.data)
    f = filing or Filing.query.filter_by(report_id=r.id).order_by(Filing.created.desc()).first()
    return dict(id=r.id, title=r.title, tool=d["tool"]["name"], created=iso(r.created), period=d["period"], narrative=r.narrative, outputs=d["outputs"],
                filing=dict(id=f.id, status=f.status, note=f.note, filing_ref=f.filing_ref) if f else None)


# ======================= .docx =======================
def build_docx(report, filing, creator, private=None):
    from docx import Document
    from docx.shared import Pt, Cm
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    ctrl = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\ufffe\uffff]")
    X = lambda s: ctrl.sub("", str(s if s is not None else ""))
    snap = json.loads(report.data)
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2)
    d.styles["Normal"].font.name, d.styles["Normal"].font.size = "Calibri", Pt(10.5)

    def shade(cell, fill):
        sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), fill)
        cell._tc.get_or_add_tcPr().append(sh)

    def kv(pairs):
        t = d.add_table(rows=0, cols=2)
        t.style = "Table Grid"
        for k, v in pairs:
            c = t.add_row().cells
            c[0].text, c[1].text = "", X(v)
            c[0].paragraphs[0].add_run(k).bold = True
            shade(c[0], "F3F5F9")
            c[0].width, c[1].width = Cm(4.2), Cm(12.8)

    d.add_heading("IP filing request: provenance report", 0)
    d.add_heading("Filing", 1)
    kv([("Protection sought", PROTECTION.get(filing.protection_type, filing.protection_type)), ("Status", STATUS.get(filing.status, "Not yet sent to an advocate")),
        ("Filing reference", filing.filing_ref or "Not yet filed"), ("Creator", "%s (%s)" % (creator.name, creator.email)),
        ("Creator's message", filing.message or "None")])
    d.add_heading("AI tool and period", 1)
    kv([("AI tool", "%s (%s)" % (snap["tool"]["name"], snap["tool"]["url"])), ("Period (UTC)", "%s to %s" % (snap["period"]["start"], snap["period"]["end"])),
        ("Visits / downloads", "%d / %d" % (snap["summary"]["visits"], snap["summary"]["downloads"]))])
    d.add_heading("Summary", 1)
    for chunk in report.narrative.split("\n\n"):
        d.add_paragraph(X(chunk))
    d.add_heading("Downloaded AI-generated content", 1)
    t = d.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for i, h in enumerate(["File", "Type / size", "Downloaded (UTC)", "SHA-256", "Link"]):
        t.rows[0].cells[i].text = h
        t.rows[0].cells[i].paragraphs[0].runs[0].bold = True
        shade(t.rows[0].cells[i], "E6F0FD")
    for o in snap["outputs"]:
        cells = t.add_row().cells
        for i, v in enumerate([o["filename"], "%s, %s" % (o["mime"] or "type unknown", "%s bytes" % o["size"] if o["size"] is not None else "size unknown"),
                               o["at"], o["sha256"] or "not recorded", o["link"] or "none"]):
            cells[i].text = ""
            cells[i].paragraphs[0].add_run(X(v)).font.size = Pt(8)
    p = d.add_paragraph()
    p.add_run(NOTE).italic = True
    if private is not None:
        d.add_heading("Advocate evaluation (private)", 1)
        d.add_paragraph(X(private or "None yet."))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def docx_response(data, name):
    return send_file(io.BytesIO(data), as_attachment=True, download_name=name,
                     mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document")


# ======================= creator API =======================
@app.get("/api/tools")
@creator_only
def api_tools():
    return jsonify([dict(key=t["key"], name=t["name"], url=t["url"]) for t in TOOLS])


@app.get("/go/<key>")
@login_required
def go(key):  # opening a tool from Manatune is logged, so a session exists even before the extension reports
    t = TOOL.get(key)
    if not t or current_user.role != "creator":
        return redirect("/")
    db.session.add(Activity(user_id=current_user.id, tool=key, kind="launch", url=t["url"], title=t["name"], created=now()))
    db.session.commit()
    return redirect(t["url"])


@app.get("/api/summary")
@creator_only
def api_summary():
    out = {t["key"]: dict(key=t["key"], name=t["name"], visits=0, downloads=0, last=None) for t in TOOLS}
    rows = Activity.query.filter_by(user_id=current_user.id).all()
    for e in rows:
        o = out.get(e.tool)
        if o and e.kind in ("visit", "download"):
            o[e.kind + "s"] += 1
        if o and (o["last"] is None or iso(e.created) > o["last"]):
            o["last"] = iso(e.created)
    return jsonify(tools=list(out.values()), extension_active=any(e.client_id for e in rows))


@app.post("/api/reports")
@creator_only
def report_create():
    tool = TOOL.get((request.get_json(silent=True) or {}).get("tool"))
    if not tool:
        return err("Choose one of the tools.", 400)
    evs = Activity.query.filter_by(user_id=current_user.id, tool=tool["key"]).order_by(Activity.created.asc()).limit(2000).all()
    if not any(e.kind == "download" for e in evs):
        return err("Nothing downloaded from this tool yet, so there is nothing to report.", 400)
    snap = snapshot(current_user, tool, evs)
    n = snap["summary"]["downloads"]
    r = Report(id=uuid.uuid4().hex, user_id=current_user.id, tool=tool["key"], narrative=summarise(snap), data=json.dumps(snap), created=now(),
               title="%s: %d download%s" % (tool["name"], n, "" if n == 1 else "s"))
    db.session.add(r)
    db.session.commit()
    return jsonify(report_json(r)), 201


@app.get("/api/reports")
@creator_only
def report_list():
    rows = Report.query.filter_by(user_id=current_user.id).order_by(Report.created.desc()).limit(100).all()
    return jsonify([report_json(r) for r in rows])


def my_report(rid):
    return Report.query.filter_by(id=rid, user_id=current_user.id).first()  # other people's ids behave like missing ones


@app.get("/api/reports/<rid>/docx")
@creator_only
def report_docx(rid):
    r = my_report(rid)
    if not r:
        return err("not found", 404)
    f = Filing.query.filter_by(report_id=rid).order_by(Filing.created.desc()).first() or Filing(protection_type="other", status="draft")
    return docx_response(build_docx(r, f, current_user), "provenance-report-%s.docx" % rid[:8])


@app.get("/api/advocates")
@creator_only
def advocates():
    return jsonify([dict(id=u.id, name=u.name) for u in User.query.filter_by(role="advocate").order_by(User.name).limit(200).all()])


@app.post("/api/filings")
@creator_only
def filing_create():
    b = request.get_json(silent=True) or {}
    r = my_report(str(b.get("report_id") or ""))
    adv = db.session.get(User, b["advocate_id"]) if isinstance(b.get("advocate_id"), int) else None
    if not r:
        return err("Choose one of your reports.", 400)
    if not adv or adv.role != "advocate":
        return err("Choose an IP advocate or firm from the list.", 400)
    if b.get("protection_type") not in PROTECTION:
        return err("Choose the kind of protection you want.", 400)
    if Filing.query.filter(Filing.report_id == r.id, Filing.status != "declined").first():
        return err("This report has already been sent to an advocate.", 409)
    f = Filing(id=uuid.uuid4().hex, report_id=r.id, creator_id=current_user.id, advocate_id=adv.id, protection_type=b["protection_type"],
               message=clean(b.get("message"), 1000), created=now())
    db.session.add(f)
    db.session.commit()
    return jsonify(ok=True, id=f.id), 201


@app.get("/api/filings")
@creator_only
def filing_list():
    return jsonify([dict(id=f.id, status=f.status) for f in Filing.query.filter_by(creator_id=current_user.id).all()])


# ======================= advocate API =======================
def adv_filing(fid):
    return Filing.query.filter_by(id=fid, advocate_id=current_user.id).first()  # other advocates' requests are 404


def request_json(f, detail=False):
    r, cu = db.session.get(Report, f.report_id), db.session.get(User, f.creator_id)
    d = dict(id=f.id, title=r.title, status=f.status, protection=PROTECTION.get(f.protection_type, f.protection_type), created=iso(f.created),
             creator_name=cu.name, creator_email=cu.email)
    if detail:
        snap = json.loads(r.data)
        d.update(message=f.message, filing_ref=f.filing_ref, evaluation=f.evaluation or "", narrative=r.narrative, outputs=snap["outputs"],
                 period=snap["period"], tool=snap["tool"]["name"], note=NOTE)
    return d


@app.get("/api/advocate/requests")
@advocate_only
def adv_list():
    return jsonify([request_json(f) for f in Filing.query.filter_by(advocate_id=current_user.id).order_by(Filing.created.desc()).limit(200).all()])


@app.get("/api/advocate/requests/<fid>")
@advocate_only
def adv_get(fid):
    f = adv_filing(fid)
    return jsonify(request_json(f, True)) if f else err("not found", 404)


@app.post("/api/advocate/requests/<fid>/decision")
@advocate_only
def adv_decide(fid):
    f = adv_filing(fid)
    if not f:
        return err("not found", 404)
    b = request.get_json(silent=True) or {}
    act, note = b.get("action"), clean(b.get("note"), 1000)
    if f.status in ("filed", "declined"):
        return err("This request is closed.", 400)
    if act == "accept" and f.status == "submitted":
        f.status, f.note = "accepted", note
    elif act == "decline":
        if not note:
            return err("Please give the creator a reason.", 400)
        f.status, f.note = "declined", note
    elif act == "evaluate" and f.status == "accepted":
        f.evaluation = str(b.get("evaluation") or "")[:5000]
    elif act == "file" and f.status == "accepted":
        ref = clean(b.get("filing_ref"), 120)
        if not ref:
            return err("Enter the filing or application reference.", 400)
        f.filing_ref, f.status, f.note = ref, "filed", note or f.note
    else:
        return err("That action isn't available for a request in this state.", 400)
    db.session.commit()
    return jsonify(request_json(f))


@app.get("/api/advocate/requests/<fid>/docx")
@advocate_only
def adv_docx(fid):
    f = adv_filing(fid)
    if not f:
        return err("not found", 404)
    return docx_response(build_docx(db.session.get(Report, f.report_id), f, db.session.get(User, f.creator_id), private=f.evaluation or ""),
                         "ip-filing-%s.docx" % f.id[:8])


# ======================= extension: tokens, download, ingestion =======================
@app.get("/api/tokens")
@creator_only
def tokens_list():
    return jsonify([dict(id=t.id, prefix=t.prefix, created=iso(t.created))
                    for t in ApiToken.query.filter_by(user_id=current_user.id, revoked=False).order_by(ApiToken.created.desc()).all()])


@app.post("/api/tokens")
@creator_only
def tokens_create():
    if ApiToken.query.filter_by(user_id=current_user.id, revoked=False).count() >= 3:
        return err("You already have 3 tokens. Revoke one first.", 400)
    raw = "mt_" + secrets.token_urlsafe(32)
    db.session.add(ApiToken(user_id=current_user.id, token_hash=hashlib.sha256(raw.encode()).hexdigest(), prefix=raw[:9], created=now()))
    db.session.commit()
    return jsonify(token=raw)


@app.delete("/api/tokens/<int:tid>")
@creator_only
def tokens_revoke(tid):
    t = ApiToken.query.filter_by(id=tid, user_id=current_user.id).first()
    if not t:
        return err("not found", 404)
    t.revoked = True
    db.session.commit()
    return jsonify(ok=True)


def _origin():
    s = urlsplit(os.environ.get("PUBLIC_URL", "").rstrip("/") or request.url_root.rstrip("/"))
    o = "%s://%s" % (s.scheme, s.netloc)
    return o if re.match(r"^https?://[A-Za-z0-9.\-]+(:\d{1,5})?$", o) else None


@app.get("/extension/download.zip")
@creator_only
def extension_zip():
    """The extension, with this site's address baked in."""
    o, folder = _origin(), os.path.join(BASE, "extension")
    if not o or not os.path.isdir(folder):
        return err("The extension is not available on this server.", 503)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name in sorted(os.listdir(folder)):
            p = os.path.join(folder, name)
            if os.path.isfile(p) and not name.startswith("."):
                data = open(p, "rb").read()
                z.writestr(name, data.replace(b"__MANATUNE_URL__", o.encode()) if name.endswith((".json", ".js")) else data)
    buf.seek(0)
    return send_file(buf, mimetype="application/zip", as_attachment=True, download_name="manatune-extension.zip")


@app.get("/api/ext/ping")
@token_required
def ext_ping():
    return jsonify(ok=True, name=g.ext_user.name)


@app.post("/api/ext/activity")
@token_required
def ext_activity():
    evs = (request.get_json(silent=True) or {}).get("events")
    if not isinstance(evs, list) or not 1 <= len(evs) <= 25:
        return err("Send between 1 and 25 events.", 400)
    uid, stored, ignored = g.ext_user.id, 0, 0
    ids = [clean(e.get("client_id"), 64) for e in evs if isinstance(e, dict)]
    have = {r[0] for r in db.session.query(Activity.client_id).filter(Activity.user_id == uid, Activity.client_id.in_([i for i in ids if i])).all()}
    for e in evs:
        if not isinstance(e, dict):
            ignored += 1
            continue
        cid, kind, url, page = clean(e.get("client_id"), 64), e.get("kind"), web_url(e.get("url")), web_url(e.get("page_url"))
        t = tool_for_url(url) or tool_for_url(page)  # anything outside the three tools is dropped
        if not t or not cid or kind not in ("visit", "download") or (kind == "visit" and not url):
            ignored += 1
            continue
        if cid in have:
            continue  # a retry of something already stored
        have.add(cid)
        dl = kind == "download"
        digest = str(e.get("sha256") or "").lower()
        size = e.get("size")
        db.session.add(Activity(
            user_id=uid, tool=t["key"], kind=kind, client_id=cid, url=url, page_url=page, title=clean(e.get("title"), 300), created=now(),
            filename=(clean(os.path.basename(str(e.get("filename") or "").replace("\\", "/")), 255) or "download") if dl else None,  # never the local folder
            mime=clean(e.get("mime"), 100) if dl else None,
            size=size if dl and isinstance(size, int) and not isinstance(size, bool) and 0 <= size <= 10 ** 12 else None,
            sha256=digest if dl and re.match(r"^[0-9a-f]{64}$", digest) else None))
        stored += 1
    try:
        db.session.commit()
    except IntegrityError:
        db.session.rollback()
        return err("Please retry.", 409)
    return jsonify(ok=True, stored=stored, ignored=ignored)


# ======================= pages + Google sign-in =======================
with app.app_context():
    for attempt in range(6):  # the database can take a moment to accept connections on boot
        try:
            db.create_all()
            break
        except OperationalError:
            if attempt == 5:
                raise
            time.sleep(3)


@app.get("/health")
def health():
    return jsonify(status="ok")


@lm.user_loader
def load_user(i):
    return db.session.get(User, int(i))


@lm.unauthorized_handler
def unauthorized():
    if request.path.startswith("/api/"):
        return jsonify(error="login required"), 401
    return redirect("/login")


def _home(tab):
    me = {"name": current_user.name, "role": current_user.role}
    if current_user.role == "advocate":
        return render_template("advocate_home.html", me=me, status=STATUS)
    return render_template("workspace.html", me=me, tab=tab, status=STATUS)


@app.get("/")
@login_required
def index():
    return _home("tools")


@app.get("/track")
@login_required
def track_page():
    return _home("track")


@app.get("/extension")
@login_required
def extension_page():
    return render_template("extension.html") if current_user.role == "creator" else redirect("/")


@app.get("/login")
def login():
    return redirect("/") if current_user.is_authenticated else render_template("login.html", mode="signin")


@app.get("/ipadvo")
def ipadvo():
    return redirect("/") if current_user.is_authenticated else render_template("login.html", mode="advocate")


@app.post("/logout")
def logout():
    back = "/ipadvo" if current_user.is_authenticated and current_user.role == "advocate" else "/login"
    logout_user()
    return redirect(back)


GOOGLE_ENABLED = bool(os.environ.get("GOOGLE_CLIENT_ID") and os.environ.get("GOOGLE_CLIENT_SECRET"))
oauth = OAuth(app)
if GOOGLE_ENABLED:
    oauth.register(name="google", client_id=os.environ["GOOGLE_CLIENT_ID"], client_secret=os.environ["GOOGLE_CLIENT_SECRET"],
                   server_metadata_url="https://accounts.google.com/.well-known/openid-configuration", client_kwargs={"scope": "openid email profile"})


def _callback_url():  # must match the Authorized redirect URI in Google Cloud exactly
    base = os.environ.get("PUBLIC_URL", "").rstrip("/")
    if base:
        return base + "/auth/google/callback"
    return url_for("google_callback", _external=True, _scheme="https" if os.environ.get("RENDER") else None)


def _start_google(intent):
    back = "/ipadvo" if intent == "advocate" else "/login"
    if not GOOGLE_ENABLED:
        flash("Google sign-in isn't configured: set GOOGLE_CLIENT_ID and GOOGLE_CLIENT_SECRET.")
        return redirect(back)
    session["intent"] = intent
    try:
        return oauth.google.authorize_redirect(_callback_url(), prompt="select_account")
    except Exception:
        app.logger.exception("Google authorize_redirect failed")
        flash("Could not reach Google. Please try again.")
        return redirect(back)


@app.get("/auth/google")
def google_login():
    return _start_google("creator")


@app.get("/auth/google/advocate")
def google_login_advocate():
    return _start_google("advocate")


@app.get("/auth/google/callback")
def google_callback():
    intent = session.pop("intent", "creator")
    back = "/ipadvo" if intent == "advocate" else "/login"
    if not GOOGLE_ENABLED or request.args.get("error"):
        flash("Google sign-in was cancelled.")
        return redirect(back)
    try:
        token = oauth.google.authorize_access_token()
        info = token.get("userinfo") or oauth.google.userinfo(token=token) or {}
    except Exception:
        app.logger.exception("Google token exchange failed")
        flash("Google sign-in failed. Please try again.")
        return redirect(back)
    email = (info.get("email") or "").strip().lower()
    if not email or info.get("email_verified") not in (True, "true"):
        flash("Google did not confirm that email address.")
        return redirect(back)
    u = User.query.filter_by(email=email).first()
    if u is None:
        name = (info.get("name") or email.split("@")[0])[:120]
        if intent == "advocate" and User.query.filter_by(role="advocate", name=name).first():
            name = ("%s (%s)" % (name, email))[:120]  # creators pick advocates by name, so names must differ
        u = User(email=email, name=name, role=intent, pw="!google-only")
        db.session.add(u)
        db.session.commit()
    elif u.role != intent:  # the two doors stay separate
        flash("This Google account is registered as an IP advocate. Use the advocate sign-in." if u.role == "advocate"
              else "This Google account is registered as a creator. Use the main sign-in.")
        return redirect("/ipadvo" if u.role == "advocate" else "/login")
    login_user(u, remember=True)
    return redirect("/")


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=False)
